#!/usr/bin/env python3
"""Generate the Sentinel-2/DJI harmonization architecture presentation."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "docs" / "sentinel2-dji-harmonization-architecture.pptx"
VERSIONED_OUTPUT = ROOT / "docs" / "geospatial-building-block-harmonization-v2.pptx"

NAVY = RGBColor(18, 36, 51)
INK = RGBColor(28, 42, 53)
MUTED = RGBColor(89, 106, 117)
PAPER = RGBColor(247, 249, 247)
WHITE = RGBColor(255, 255, 255)
TEAL = RGBColor(0, 122, 116)
TEAL_LIGHT = RGBColor(214, 239, 235)
CORAL = RGBColor(221, 89, 67)
CORAL_LIGHT = RGBColor(250, 226, 220)
GOLD = RGBColor(224, 164, 46)
GOLD_LIGHT = RGBColor(250, 239, 211)
BLUE = RGBColor(45, 104, 160)
BLUE_LIGHT = RGBColor(220, 233, 244)
GREEN = RGBColor(51, 132, 88)
RED = RGBColor(178, 61, 52)

SLIDE_W = 13.333
SLIDE_H = 7.5


def set_background(slide, color=PAPER):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_text(
    slide,
    text,
    left,
    top,
    width,
    height,
    size=18,
    color=INK,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.06)
    frame.margin_right = Inches(0.06)
    frame.margin_top = Inches(0.03)
    frame.margin_bottom = Inches(0.03)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_box(
    slide,
    text,
    left,
    top,
    width,
    height,
    fill,
    line,
    size=16,
    color=INK,
    bold=True,
):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1.2)
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.08)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def add_arrow(slide, left, top, width=0.42, height=0.35, color=MUTED):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.CHEVRON,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_header(slide, title, kicker=None):
    if kicker:
        add_text(slide, kicker.upper(), 0.55, 0.20, 3.0, 0.28, 10, TEAL, True)
    add_text(slide, title, 0.55, 0.48, 12.1, 0.58, 27, NAVY, True)
    rule = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(0.55), Inches(1.09), Inches(12.2), Inches(0.035)
    )
    rule.fill.solid()
    rule.fill.fore_color.rgb = TEAL
    rule.line.fill.background()


def add_footer(slide, number):
    add_text(
        slide,
        "Geospatial Building Block | Sentinel-2 + DJI",
        0.55,
        7.12,
        5.0,
        0.2,
        9,
        MUTED,
    )
    add_text(slide, str(number), 12.25, 7.10, 0.5, 0.2, 9, MUTED, align=PP_ALIGN.RIGHT)


def add_bullets(slide, items, left, top, width, height, size=17, color=INK):
    box = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = color
        paragraph.space_after = Pt(10)
        paragraph.text = f"• {item}"
    return box


def new_slide(prs, title, number, kicker=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, title, kicker)
    add_footer(slide, number)
    return slide


def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, NAVY)
    accent = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(0), Inches(0), Inches(0.18), Inches(SLIDE_H)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = CORAL
    accent.line.fill.background()
    add_text(slide, "GEOSPATIAL BUILDING BLOCK | HOW IT WORKS", 0.72, 0.72, 7.5, 0.32, 12, GOLD, True)
    add_text(
        slide,
        "From process contract\nto comparable imagery",
        0.68,
        1.45,
        8.8,
        1.75,
        38,
        WHITE,
        True,
    )
    add_text(
        slide,
        "Kernel architecture, execution lifecycle, and a runnable\nSentinel-2 + DJI harmonization example",
        0.72,
        3.48,
        8.4,
        0.9,
        19,
        RGBColor(204, 218, 225),
    )
    add_box(slide, "SENTINEL-2\nREFERENCE", 9.65, 1.25, 2.55, 1.2, BLUE, BLUE, 16, WHITE)
    add_arrow(slide, 10.70, 2.75, 0.45, 0.55, GOLD)
    add_box(slide, "DJI\nCOMPARISON", 9.65, 3.62, 2.55, 1.2, CORAL, CORAL, 16, WHITE)
    add_text(slide, "Revision 2 | 15 September 2026 | Architecture, contracts, execution, and evidence", 0.72, 6.55, 8.6, 0.35, 13, RGBColor(157, 179, 190))

    slide = new_slide(prs, "What is a Geospatial Building Block?", 2, "Reusable capability with observable behavior")
    add_box(slide, "INPUT PROFILE", 0.55, 1.55, 2.28, 0.62, BLUE, BLUE, 15, WHITE)
    add_text(slide, "STAC, GeoTIFF, GeoJSON, or a mapped custom profile", 0.62, 2.34, 2.15, 1.15, 15, MUTED, False, PP_ALIGN.CENTER)
    add_arrow(slide, 3.00, 2.02, 0.55, 0.46, TEAL)
    add_box(slide, "ALGORITHM CONTRACT", 3.75, 1.55, 2.55, 0.62, TEAL, TEAL, 15, WHITE)
    add_text(slide, "Process identity, accepted semantics, parameters, runtime", 3.88, 2.34, 2.30, 1.15, 15, MUTED, False, PP_ALIGN.CENTER)
    add_arrow(slide, 6.48, 2.02, 0.55, 0.46, TEAL)
    add_box(slide, "TYPED OUTPUT", 7.22, 1.55, 2.28, 0.62, GOLD, GOLD, 15, WHITE)
    add_text(slide, "Declared identity, spatial/temporal scope, media type", 7.34, 2.34, 2.02, 1.15, 15, MUTED, False, PP_ALIGN.CENTER)
    add_arrow(slide, 9.68, 2.02, 0.55, 0.46, CORAL)
    add_box(slide, "EVIDENCE", 10.42, 1.55, 2.28, 0.62, CORAL, CORAL, 15, WHITE)
    add_text(slide, "Execution ID, input/output references, validation result", 10.54, 2.34, 2.02, 1.15, 15, MUTED, False, PP_ALIGN.CENTER)
    add_box(slide, "Metadata are context-dependent. The kernel dynamically creates contracts that bind the relevant data and process context, then harmonizes them through shared semantics, behavior, and evidence.", 1.05, 4.18, 11.18, 0.92, TEAL_LIGHT, TEAL, 16, NAVY)
    add_bullets(slide, ["No fixed metadata schema is sufficient for every dynamic workflow", "Discoverable through contracts: logical representations of data, processes, and their context", "Context-specific metadata remain useful inside the declared contract", "Executable and auditable through shared framework rules"], 1.38, 5.30, 10.6, 1.42, 14)

    slide = new_slide(prs, "Functional contract and interface bindings", 3, "The kernel has four cooperating layers")
    layers = [
        ("INTERFACE BINDINGS", "REST | OGC API - Processes | OpenAPI | openEO | Records", BLUE, BLUE_LIGHT),
        ("FUNCTIONAL CONTRACT", "Adapt input profile -> validate -> execute -> publish typed output", TEAL, TEAL_LIGHT),
        ("PROCESS IMPLEMENTATION", "CWL graph + Python/GDAL/Rasterio/TorchScript + runtime environment", GOLD, GOLD_LIGHT),
        ("SEMANTICS + EVIDENCE", "Process-Type Register | STAC | quality report | checksums | W3C PROV", CORAL, CORAL_LIGHT),
    ]
    for index, (label, detail, line, fill) in enumerate(layers):
        y = 1.38 + index * 1.30
        add_box(slide, label, 0.72, y, 2.80, 0.78, line, line, 14, WHITE)
        add_arrow(slide, 3.72, y + 0.21, 0.48, 0.36, line)
        add_box(slide, detail, 4.38, y, 8.25, 0.78, fill, line, 15, NAVY, False)
    add_text(slide, "Interface choice can vary per Building Block; the evidence contract remains stable.", 1.2, 6.66, 10.9, 0.30, 14, NAVY, True, PP_ALIGN.CENTER)

    slide = new_slide(prs, "How a request moves through the Building Block", 4, "Execution lifecycle")
    lifecycle = [
        ("1", "DISCOVER", "Read kernel metadata and process semantics", BLUE),
        ("2", "ADAPT", "Map STAC/raster inputs to the algorithm contract", TEAL),
        ("3", "EXECUTE", "Run CWL locally or submit the job to REANA", GOLD),
        ("4", "VALIDATE", "Apply grid, quality, and policy acceptance gates", CORAL),
        ("5", "PUBLISH", "Return typed assets, STAC metadata, and provenance", GREEN),
    ]
    for index, (num, label, detail, color) in enumerate(lifecycle):
        x = 0.42 + index * 2.58
        add_box(slide, num, x + 0.72, 1.45, 0.58, 0.58, color, color, 16, WHITE)
        add_text(slide, label, x, 2.24, 2.05, 0.40, 17, NAVY, True, PP_ALIGN.CENTER)
        add_text(slide, detail, x, 2.82, 2.05, 1.30, 14, MUTED, False, PP_ALIGN.CENTER)
        if index < 4:
            add_arrow(slide, x + 2.10, 1.58, 0.36, 0.34, color)
    add_box(slide, "Failure is also a typed result: invalid inputs or failed thresholds stop publication and preserve diagnostic evidence.", 1.00, 5.12, 11.30, 0.92, CORAL_LIGHT, CORAL, 16, NAVY)

    slide = new_slide(prs, "Example: two workflows, one shared contract", 5, "Imagery harmonization Building Block")
    add_box(slide, "Sentinel-2\nL2A RGB", 0.45, 1.65, 1.55, 0.82, BLUE_LIGHT, BLUE, 15)
    add_arrow(slide, 2.10, 1.88)
    add_box(slide, "Co-register +\nreproject", 2.62, 1.65, 1.72, 0.82, BLUE_LIGHT, BLUE, 14)
    add_arrow(slide, 4.45, 1.88)
    add_box(slide, "4x AI super-\nresolution", 4.97, 1.65, 1.82, 0.82, BLUE, BLUE, 14, WHITE)
    add_arrow(slide, 6.90, 1.88)
    add_box(slide, "Reference image\n+ histogram", 7.42, 1.65, 2.00, 0.82, TEAL_LIGHT, TEAL, 14)

    add_box(slide, "DJI frames +\nposes + DSM", 0.45, 4.05, 1.55, 0.82, CORAL_LIGHT, CORAL, 14)
    add_arrow(slide, 2.10, 4.28)
    add_box(slide, "Photogrammetric\northomosaic", 2.62, 4.05, 1.72, 0.82, CORAL_LIGHT, CORAL, 14)
    add_arrow(slide, 4.45, 4.28)
    add_box(slide, "Area\ndownsampling", 4.97, 4.05, 1.82, 0.82, CORAL, CORAL, 14, WHITE)
    add_arrow(slide, 6.90, 4.28)
    add_box(slide, "Masked histogram\nmatching", 7.42, 4.05, 2.00, 0.82, GOLD_LIGHT, GOLD, 14)

    add_box(slide, "COMMON GRID + QUALITY GATE", 10.05, 2.60, 2.72, 1.35, TEAL, TEAL, 16, WHITE)
    add_arrow(slide, 9.48, 2.65, 0.45, 0.35, TEAL)
    add_arrow(slide, 9.48, 3.75, 0.45, 0.35, TEAL)
    add_text(slide, "COGs | STAC Items | W3C PROV | QC metrics", 3.5, 5.65, 6.8, 0.4, 16, NAVY, True, PP_ALIGN.CENTER)

    slide = new_slide(prs, "Workflow 1: Sentinel-2 establishes the reference", 6, "Executable processing")
    stages = [
        ("01", "Validate", "L2A level, RGB bands, cloud and nodata masks"),
        ("02", "Align", "Reproject and co-register to AOI and target CRS"),
        ("03", "Enhance", "TorchScript 4x inference or labelled bicubic demo"),
        ("04", "Profile", "256-bin masked histograms and robust quantiles"),
        ("05", "Publish", "COG, target grid, QC, STAC, and provenance"),
    ]
    for index, (num, title, detail) in enumerate(stages):
        x = 0.48 + index * 2.56
        add_box(slide, num, x, 1.48, 0.56, 0.56, BLUE, BLUE, 15, WHITE)
        add_text(slide, title, x, 2.15, 2.15, 0.38, 18, NAVY, True)
        add_text(slide, detail, x, 2.62, 2.18, 1.15, 14, MUTED)
        if index < len(stages) - 1:
            add_arrow(slide, x + 2.18, 1.58, 0.30, 0.32, BLUE)
    add_box(slide, "Interpretation product, not new optical truth", 0.68, 4.32, 5.55, 0.72, GOLD_LIGHT, GOLD, 17, NAVY)
    add_bullets(
        slide,
        [
            "Preserve the original 10 m source link",
            "Validate tensor shape, scale, target resolution, and model checksum",
            "Apply cloud/nodata masks during inference and histogram sampling",
        ],
        6.60,
        4.12,
        5.95,
        1.85,
        15,
    )

    slide = new_slide(prs, "Workflow 2: DJI conforms to the reference", 7, "Photogrammetry handoff + executable processing")
    add_box(slide, "RAW INPUTS", 0.62, 1.45, 2.1, 0.52, CORAL, CORAL, 14, WHITE)
    add_bullets(slide, ["Images", "Calibration", "RTK/PPK or poses", "DEM/DSM", "Optional GCPs"], 0.68, 2.08, 2.6, 2.65, 15)
    add_arrow(slide, 3.18, 3.00, 0.58, 0.46, CORAL)
    add_box(slide, "PHOTOGRAMMETRY", 4.00, 1.45, 2.25, 0.52, CORAL, CORAL, 14, WHITE)
    add_bullets(slide, ["Camera model", "Tie points", "Dense surface", "Orthomosaic", "Checkpoint RMSE"], 4.06, 2.08, 2.65, 2.65, 15)
    add_arrow(slide, 6.73, 3.00, 0.58, 0.46, CORAL)
    add_box(slide, "HARMONIZATION", 7.55, 1.45, 2.25, 0.52, GOLD, GOLD, 14, WHITE)
    add_bullets(slide, ["Exact target grid", "Area/average kernel", "Antialiasing", "Quantile matching", "Mask preservation"], 7.61, 2.08, 2.65, 2.65, 15)
    add_arrow(slide, 10.25, 3.00, 0.58, 0.46, TEAL)
    add_box(slide, "COMPARABLE\nDJI COG", 11.02, 2.55, 1.72, 1.28, TEAL, TEAL, 16, WHITE)
    add_box(slide, "Why no AI downsampler? Area aggregation is deterministic, spectrally traceable, and easier to validate.", 2.15, 5.52, 9.05, 0.72, TEAL_LIGHT, TEAL, 15, NAVY)

    slide = new_slide(prs, "The handoff is a machine-readable contract", 8, "What makes the outputs comparable")
    left_items = [
        "Same CRS and affine transform",
        "Same width, height, extent, and 2.5 m grid",
        "Same RGB band order, dtype, range, and nodata",
        "Separate cloud, shadow, saturation, and validity masks",
    ]
    right_items = [
        "Reference profile uses valid common overlap",
        "Per-band clipped quantile mapping",
        "Before/after histogram distances retained",
        "Time, sun angle, and view geometry remain explicit",
    ]
    add_box(slide, "SPATIAL CONTRACT", 0.62, 1.45, 5.78, 0.58, BLUE, BLUE, 16, WHITE)
    add_bullets(slide, left_items, 0.72, 2.18, 5.55, 2.85, 16)
    add_box(slide, "RADIOMETRIC CONTRACT", 6.92, 1.45, 5.78, 0.58, GOLD, GOLD, 16, WHITE)
    add_bullets(slide, right_items, 7.02, 2.18, 5.55, 2.85, 16)
    add_box(slide, "Histogram similarity supports interpretation; it does not erase seasonal, atmospheric, or sensor differences.", 1.18, 5.55, 10.98, 0.78, CORAL_LIGHT, CORAL, 16, NAVY)

    slide = new_slide(prs, "Quality gate: comparability is measured", 9, "Example acceptance criteria")
    checks = [
        ("Pixel grid", "Exact equality", GREEN),
        ("Registration RMSE", "< 1 target pixel", GREEN),
        ("Common overlap", ">= 10,000 pixels", GREEN),
        ("Histogram distance", "JS <= 0.10 / band", GOLD),
        ("Clipping", "<= 1%", GOLD),
        ("Cloud + shadow", "<= 5%", GOLD),
    ]
    for index, (name, threshold, color) in enumerate(checks):
        row = index // 3
        col = index % 3
        x = 0.60 + col * 4.22
        y = 1.55 + row * 2.25
        add_box(slide, name.upper(), x, y, 3.75, 0.52, color, color, 13, WHITE)
        add_text(slide, threshold, x, y + 0.70, 3.75, 0.55, 21, NAVY, True, PP_ALIGN.CENTER)
    add_text(slide, "Thresholds are deployment parameters, not universal accuracy claims.", 2.15, 6.30, 9.05, 0.36, 14, RED, True, PP_ALIGN.CENTER)

    slide = new_slide(prs, "What the execution produces", 10, "Evidence accompanies every derived image")
    evidence = [
        ("RASTER", "Cloud Optimized GeoTIFF\nExact target grid + overviews", BLUE, BLUE_LIGHT),
        ("CATALOG", "STAC Item\nGeometry + projection + asset checksum", TEAL, TEAL_LIGHT),
        ("QUALITY", "Pixel overlap + RMSE\nJS distance + clipping + masks", GOLD, GOLD_LIGHT),
        ("PROVENANCE", "Input/output SHA-256\nMethod + model identity + run ID", CORAL, CORAL_LIGHT),
    ]
    for index, (label, detail, color, fill) in enumerate(evidence):
        x = 0.52 + index * 3.18
        add_box(slide, label, x, 1.52, 2.72, 0.55, color, color, 14, WHITE)
        add_box(slide, detail, x, 2.25, 2.72, 1.42, fill, color, 15, NAVY, False)
    add_box(slide, "Fixture result", 0.85, 4.45, 2.12, 0.58, NAVY, NAVY, 15, WHITE)
    add_text(slide, "65,536", 3.34, 4.36, 1.55, 0.52, 24, BLUE, True, PP_ALIGN.CENTER)
    add_text(slide, "overlap pixels", 3.23, 4.91, 1.78, 0.30, 12, MUTED, False, PP_ALIGN.CENTER)
    add_text(slide, "0.138 -> 0.0043", 5.32, 4.36, 2.35, 0.52, 24, TEAL, True, PP_ALIGN.CENTER)
    add_text(slide, "worst-band JS distance", 5.40, 4.91, 2.18, 0.30, 12, MUTED, False, PP_ALIGN.CENTER)
    add_text(slide, "0%", 8.13, 4.36, 1.30, 0.52, 24, GOLD, True, PP_ALIGN.CENTER)
    add_text(slide, "clipping", 8.13, 4.91, 1.30, 0.30, 12, MUTED, False, PP_ALIGN.CENTER)
    add_text(slide, "PASS", 10.12, 4.36, 1.62, 0.52, 24, GREEN, True, PP_ALIGN.CENTER)
    add_text(slide, "all gates", 10.12, 4.91, 1.62, 0.30, 12, MUTED, False, PP_ALIGN.CENTER)
    add_text(slide, "A deliberately tightened JS threshold returns exit code 2 and a machine-readable failed report.", 1.0, 6.10, 11.3, 0.42, 14, RED, True, PP_ALIGN.CENTER)

    slide = new_slide(prs, "How it executes today", 11, "CWL locally; REANA packaging supplied")
    add_box(slide, "1. RUN SENTINEL-2", 0.62, 1.48, 3.65, 0.58, BLUE, BLUE, 15, WHITE)
    add_text(slide, "Produces target_grid.json and\nreference_histogram.json", 0.72, 2.23, 3.45, 0.82, 17, NAVY, True, PP_ALIGN.CENTER)
    add_arrow(slide, 4.55, 2.10, 0.62, 0.48, TEAL)
    add_box(slide, "2. TRANSFER CONTRACTS", 5.43, 1.48, 3.65, 0.58, TEAL, TEAL, 15, WHITE)
    add_text(slide, "Checksum and upload to the\nsecond REANA run", 5.53, 2.23, 3.45, 0.82, 17, NAVY, True, PP_ALIGN.CENTER)
    add_arrow(slide, 9.35, 2.10, 0.62, 0.48, CORAL)
    add_box(slide, "3. RUN DJI", 10.23, 1.48, 2.48, 0.58, CORAL, CORAL, 15, WHITE)
    add_text(slide, "Produces matched COG\nand QC evidence", 10.28, 2.23, 2.38, 0.82, 16, NAVY, True, PP_ALIGN.CENTER)
    add_box(slide, "LOCAL", 0.78, 4.08, 2.32, 0.55, GREEN, GREEN, 14, WHITE)
    add_bullets(slide, ["cwltool executes now", "run_demo.py chains both runs", "CWL stages scripts explicitly", "Failure gates return nonzero"], 0.88, 4.82, 3.55, 1.50, 14)
    add_box(slide, "REANA", 5.02, 4.08, 2.32, 0.55, TEAL, TEAL, 14, WHITE)
    add_bullets(slide, ["Two descriptors supplied", "Reference assets transfer between runs", "Worker image must include dependencies", "Runtime logs augment provenance"], 5.12, 4.82, 3.55, 1.50, 14)
    add_box(slide, "API FACADE", 9.22, 4.08, 2.32, 0.55, CORAL, CORAL, 14, WHITE)
    add_bullets(slide, ["Kernel metadata is registered", "Existing facade proves job pattern", "Harmonization endpoints remain next", "Bindings are target capabilities"], 9.32, 4.82, 3.35, 1.50, 14)
    add_text(slide, "Current API boundary: direct CWL execution works; dedicated harmonization facade endpoints are not yet implemented.", 1.25, 6.53, 10.85, 0.32, 13, RED, True, PP_ALIGN.CENTER)

    slide = new_slide(prs, "Executable package and production boundary", 12, "What exists and what remains")
    add_box(slide, "RUN NOW", 0.62, 1.40, 5.82, 0.58, GREEN, GREEN, 15, WHITE)
    add_bullets(slide, [
        "workflows/imagery_harmonization/run_demo.py",
        "sentinel2-workflow.cwl + dji-workflow.cwl",
        "Rasterio COG/grid/histogram processing",
        "TorchScript interface and model provenance",
        "Positive and failure acceptance behavior",
    ], 0.76, 2.18, 5.50, 3.35, 15)
    add_box(slide, "PRODUCTION INTEGRATION", 6.88, 1.40, 5.82, 0.58, CORAL, CORAL, 15, WHITE)
    add_bullets(slide, [
        "Validated Sentinel-2 model + completed model card",
        "ODM/WebODM orthomosaic and checkpoint report",
        "Tiled model inference for large scenes",
        "Published REANA worker image",
        "Facade process endpoints and field calibration",
    ], 7.02, 2.18, 5.50, 3.35, 15)
    add_box(slide, "python3 workflows/imagery_harmonization/run_demo.py", 2.05, 6.08, 9.20, 0.62, NAVY, NAVY, 17, WHITE)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    prs.save(VERSIONED_OUTPUT)
    print(OUTPUT)
    print(VERSIONED_OUTPUT)


if __name__ == "__main__":
    build_deck()